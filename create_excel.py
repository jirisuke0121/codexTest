import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

# シート1: 振り返り内容一覧表
data1 = {
    '大項目': [
        '(1) 参加動機に対する気づき',
        '(1) 参加動機に対する気づき',
        '(2) My目標（技術）',
        '(2) My目標（技術）',
        '(3) My目標（職業準備性）',
        '(3) My目標（職業準備性）'
    ],
    '評価': [4, 4, 3, 3, 4, 4],
    '小項目': [
        '① 上流工程・コミュニケーション面の課題',
        '② データ分析PJへの挑戦',
        '① アンケートデータ分析・可視化の実践',
        '② 分析フロー設計',
        '① チーム内での立ち回り',
        '② 判断の癖・傾向の言語化'
    ],
    '記載内容': [
        '現職では上流工程を担う中で、PJ上流での認識合わせや役割認識、期待値調整に課題を感じていた。疑似就労ではPLとして体制作りや役割分担、顧客との合意形成を経験し、前提条件や制約を明示した上で進めること、判断や収束を誰が引き取るかを明確にすることの重要性を再認識した。',
        'アンケートデータを扱い、不確実なデータ状況下でも仮説検証の方向性を示し、関係者と合意形成を行う経験ができた。分析結果そのもの以上に、「どこまで分かっていて、どこからが不確実か」を整理して説明するプロセスに実務的な価値があると理解でき、今後のキャリア検討においても重要な経験となった。',
        'Pythonを用いて、見学アンケート・体験アンケート・契約KPIを横断的に整理・分析した。その結果、「体験→契約フェーズが最も大きなボトルネックであること」や「利用者の不満が一部項目に集中している可能性」などを示すことができた。定量的な指標と定性的な記述を組み合わせることで、アンケート改善を優先すべき理由を、関係者へ説明できる状態まで整理できたと考えている。',
        '分析にあたっては、前処理の洗い出しや分析対象の絞り込みを明示的に行い、属人的になりにくい進め方を意識した。「どう効果的に分析を進めるか」という思考プロセスの型を意識できた点が成果と感じている。',
        'PLとして進捗状況や負荷を踏まえながら役割分担の見直しや優先順位調整を行った。議論の場では、自身の理解内容を整理して共有し、チーム内での認識合わせを促す役割を担う場面が多かった。また、SE（兼データエンジニア）として、顧客視点で求められている成果や制約条件を意識しながら、資料作成・説明・プレゼンテーションに取り組んだ。チーム内では役割や期限を意識しつつメンバーの意見を引き出し、最終的にはPLとして判断・収束を引き取る立場を経験できたことは、今後の職業生活に向けた準備として大きな学びとなった。',
        'PLの立場になることで全体像が見える分、メンバーから出てくるアウトプットが、PLとして期待する水準に達していないと感じる場面があった。その結果、PLが先回りしてタスクを巻き取り、自身のアウトプットで成立させてしまう判断を取りがちだと感じた。しかし、この進め方は短期的にはPJを前に進められる一方で、メンバーが「何が期待されているアウトプットなのか」を学ぶ機会を失い、参加意識や当事者意識が育たないまま、PLに負荷が集中し続ける構造を生みやすい。今回の振り返りを通じて、短期的なアウトプットよりも、チームとして最終的に成果が成立することを優先しがちであること、PLが抱え込みすぎることで負荷の偏りや参加意識の低下につながりやすい構造があることに気づいた。今後は、この点を意識的な判断ポイントとして活かしていきたいと考えている。'
    ]
}
df1 = pd.DataFrame(data1)

# シート2: 項目別サマリー表
data2 = {
    '大項目': [
        '(1) 参加動機',
        '(2) 技術目標',
        '(3) 職業準備性'
    ],
    '評価': [4, 3, 4],
    '主な経験・取り組み': [
        'PLとして体制作り・役割分担・顧客との合意形成／不確実なデータ状況下での仮説検証・合意形成',
        'Pythonでアンケート・KPIを横断分析／前処理・分析対象の絞り込みを明示化',
        'PLとして役割分担・優先順位調整／SE兼データエンジニアとして資料作成・プレゼン／メンバーの意見を引き出し判断・収束を引き取る'
    ],
    '気づき・学び': [
        '前提条件・制約の明示と判断者の明確化の重要性／「確定・不確定の境界を整理・説明するプロセス」の実務的価値',
        'ボトルネック特定・不満集中傾向の可視化／属人化回避の進め方・分析の思考プロセスの型',
        'タスクを巻き取る傾向の認識／メンバーの学習機会喪失・当事者意識低下・PL負荷集中の構造への気づき'
    ]
}
df2 = pd.DataFrame(data2)

# シート3: 評価観点別整理表
data3 = {
    '評価観点': [
        '技術', '技術', '技術',
        '職業準備性', '職業準備性', '職業準備性',
        'PL視点', 'PL視点', 'PL視点', 'PL視点'
    ],
    '該当箇所': [
        '(2)-①', '(2)-②', '(1)-②',
        '(3)-①', '(3)-①', '(3)-②',
        '(1)-①', '(1)-①', '(3)-①', '(3)-②'
    ],
    '内容': [
        'Pythonによるアンケートデータ分析・可視化（見学・体験・契約KPI横断）',
        '分析フロー設計（前処理洗い出し・分析対象絞り込みの明示化）',
        '不確実なデータ状況下での仮説検証',
        'チーム内での立ち回り（役割分担・優先順位調整・認識合わせ促進）',
        '資料作成・説明・プレゼンテーション',
        '自身の判断の癖・傾向の言語化',
        '体制作り・役割分担・顧客との合意形成',
        '前提条件・制約の明示、判断者の明確化',
        'メンバーの意見を引き出し、最終判断・収束を引き取る',
        'タスク巻き取り傾向とその弊害（負荷集中・当事者意識低下）への気づき'
    ]
}
df3 = pd.DataFrame(data3)

# Excelファイル作成
wb = Workbook()

# スタイル定義
header_font = Font(bold=True, size=11)
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font_white = Font(bold=True, size=11, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)
wrap_alignment = Alignment(wrap_text=True, vertical='top')

def style_sheet(ws, df, col_widths):
    # ヘッダー行のスタイル
    for col_idx, cell in enumerate(ws[1], 1):
        cell.font = header_font_white
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # データ行のスタイル
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), 2):
        for cell in row:
            cell.border = thin_border
            cell.alignment = wrap_alignment
    
    # 列幅設定
    for col_idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + col_idx)].width = width

# シート1: 振り返り内容一覧表
ws1 = wb.active
ws1.title = "振り返り内容一覧"
for r in dataframe_to_rows(df1, index=False, header=True):
    ws1.append(r)
style_sheet(ws1, df1, [30, 8, 35, 80])
ws1.row_dimensions[1].height = 25
for row_idx in range(2, ws1.max_row + 1):
    ws1.row_dimensions[row_idx].height = 80

# シート2: 項目別サマリー表
ws2 = wb.create_sheet(title="項目別サマリー")
for r in dataframe_to_rows(df2, index=False, header=True):
    ws2.append(r)
style_sheet(ws2, df2, [20, 8, 50, 50])
ws2.row_dimensions[1].height = 25
for row_idx in range(2, ws2.max_row + 1):
    ws2.row_dimensions[row_idx].height = 60

# シート3: 評価観点別整理表
ws3 = wb.create_sheet(title="評価観点別整理")
for r in dataframe_to_rows(df3, index=False, header=True):
    ws3.append(r)
style_sheet(ws3, df3, [15, 12, 70])
ws3.row_dimensions[1].height = 25
for row_idx in range(2, ws3.max_row + 1):
    ws3.row_dimensions[row_idx].height = 30

# 保存
wb.save('/workspace/振り返り整理表.xlsx')
print("Excelファイルを作成しました: /workspace/振り返り整理表.xlsx")
